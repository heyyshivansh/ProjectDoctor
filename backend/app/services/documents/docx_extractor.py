import os
from typing import List
import docx
from docx.opc.exceptions import PackageNotFoundError
from app.services.documents.base import BaseDocumentExtractor, ExtractionResult, DocumentSection

MAX_DOCX_FILE_SIZE_BYTES = 50 * 1024 * 1024  # 50MB ceiling for docx


class DocxExtractor(BaseDocumentExtractor):
    """Deterministic Word Document (.docx) text extractor using python-docx."""

    def extract(self, file_path: str) -> ExtractionResult:
        if not os.path.exists(file_path):
            return ExtractionResult(
                status="failed",
                extractor_name="python-docx",
                error_message=f"File not found: {file_path}",
            )

        # Defensive file size check before unzipping
        file_size = os.path.getsize(file_path)
        if file_size > MAX_DOCX_FILE_SIZE_BYTES:
            return ExtractionResult(
                status="failed",
                extractor_name="python-docx",
                error_message=f"DOCX file size ({file_size} bytes) exceeds safety ceiling of {MAX_DOCX_FILE_SIZE_BYTES} bytes.",
            )

        try:
            doc = docx.Document(file_path)
        except PackageNotFoundError:
            return ExtractionResult(
                status="failed",
                extractor_name="python-docx",
                error_message="Invalid or corrupted DOCX package.",
            )
        except Exception as e:
            return ExtractionResult(
                status="failed",
                extractor_name="python-docx",
                error_message=f"Failed to open DOCX document: {str(e)}",
            )

        extracted_parts: List[str] = []
        sections: List[DocumentSection] = []
        current_char_offset = 0

        # Extract paragraphs and detect heading styles
        for para in doc.paragraphs:
            text = para.text
            if not text.strip():
                continue

            style_name = para.style.name if para.style else ""
            is_heading = style_name.startswith("Heading") or style_name.startswith("Title")
            level = 1
            if "2" in style_name:
                level = 2
            elif "3" in style_name:
                level = 3
            elif "4" in style_name:
                level = 4

            formatted_para = text + "\n\n"
            start_char = current_char_offset
            end_char = start_char + len(formatted_para)

            if is_heading:
                sections.append(
                    DocumentSection(
                        title=text.strip()[:150],
                        level=level,
                        start_char=start_char,
                        end_char=end_char,
                        text_preview=text.strip()[:100],
                    )
                )

            extracted_parts.append(formatted_para)
            current_char_offset = end_char

        # Extract table text
        for table_idx, table in enumerate(doc.tables):
            table_header = f"--- Table {table_idx + 1} ---\n"
            table_lines: List[str] = [table_header]
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_lines.append(" | ".join(row_cells))
            
            table_content = "\n".join(table_lines) + "\n\n"
            start_char = current_char_offset
            end_char = start_char + len(table_content)

            sections.append(
                DocumentSection(
                    title=f"Table {table_idx + 1}",
                    level=2,
                    start_char=start_char,
                    end_char=end_char,
                    text_preview=table_content.strip()[:100],
                )
            )
            extracted_parts.append(table_content)
            current_char_offset = end_char

        raw_text = "".join(extracted_parts)

        # Docx core properties metadata
        doc_metadata = {}
        try:
            props = doc.core_properties
            if props.title:
                doc_metadata["title"] = props.title
            if props.author:
                doc_metadata["author"] = props.author
            if props.comments:
                doc_metadata["comments"] = props.comments
        except Exception:
            pass

        doc_metadata["paragraph_count"] = len(doc.paragraphs)
        doc_metadata["table_count"] = len(doc.tables)

        return ExtractionResult(
            status="completed",
            extractor_name="python-docx",
            raw_text=raw_text,
            page_count=None,
            sections=sections,
            metadata=doc_metadata,
        )
